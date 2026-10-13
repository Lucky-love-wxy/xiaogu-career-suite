#!/usr/bin/env python3
"""把职业库导出成人工审查用的 Excel：每个有样本的岗位一行，附样本明细和专业连线，黄色列留给审查人填。

数字和原文全部来自 data/careerdb（cards / edges / postings_mapped / role_cards / stats），不生成新事实。
单位数在这里按 postings_mapped 里挂到该节点的不同 company 现算，因为 cards.json 没有这个字段。
"""
import json
from collections import Counter, defaultdict
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[2]
DB = ROOT / "data" / "careerdb"
OUT = ROOT / "docs" / "research" / "careerdb-review-2026-10-07.xlsx"
SEP = " › "

FONT = Font(name="Arial", size=10)
HEAD = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor="44546A")
INPUT_FILL = PatternFill("solid", fgColor="FFFF00")  # 审查人填写列
INPUT_HEAD_FILL = PatternFill("solid", fgColor="BF9000")
WRAP = Alignment(wrap_text=True, vertical="top")

ROLE_VERDICTS = ["通过", "归类有误", "应拆分", "应合并", "证据不足", "删除", "待定"]
SAMPLE_VERDICTS = ["正确", "归错节点", "应新建节点", "不是岗位/无效样本"]


def counts(d, k=None):
    return "、".join(f"{a}{b}" for a, b in sorted(d.items(), key=lambda x: -x[1])[:k])


def write_table(ws, headers, rows, widths, inputs=(), verdicts=None):
    """写表头+数据，冻结首行、加筛选；inputs 是审查人填写的列名，标黄并给第一列下拉。"""
    ws.append(headers)
    for r in rows:
        ws.append(r)
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=i)
        c.font = HEAD
        c.fill = INPUT_HEAD_FILL if h in inputs else HEAD_FILL
        c.alignment = Alignment(wrap_text=True, vertical="center")
        ws.column_dimensions[get_column_letter(i)].width = widths.get(h, 12)
    n = len(rows) + 1
    for row in ws.iter_rows(min_row=2, max_row=n):
        for c in row:
            c.font = FONT
            c.alignment = WRAP
            if headers[c.column - 1] in inputs:
                c.fill = INPUT_FILL
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{n}"
    if verdicts and inputs:
        col = get_column_letter(headers.index(inputs[0]) + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(verdicts) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        dv.add(f"{col}2:{col}{n}")


def link(ws, header, headers, n):
    """把链接列变成可点的超链接（文本保留原 URL）。"""
    col = headers.index(header) + 1
    for r in range(2, n + 1):
        c = ws.cell(row=r, column=col)
        if c.value and str(c.value).startswith("http"):
            c.hyperlink = c.value
            c.font = Font(name="Arial", size=10, color="0563C1", underline="single")


def main():
    cards = json.loads((DB / "cards.json").read_text(encoding="utf-8"))
    edges = json.loads((DB / "edges.json").read_text(encoding="utf-8"))
    stats = json.loads((DB / "stats.json").read_text(encoding="utf-8"))
    posts = [json.loads(l) for l in (DB / "postings_mapped.jsonl").read_text(encoding="utf-8").splitlines() if l]
    rc = {}
    for f in (DB / "role_cards").glob("*.json"):
        d = json.loads(f.read_text(encoding="utf-8"))
        rc[d["node"]] = d

    by_node = defaultdict(list)
    for p in posts:
        if p["node"]:
            by_node[p["node"]].append(p)

    wb = Workbook()
    readme = wb.active
    readme.title = "说明"

    # 岗位审查：每个有样本的节点一行，按样本数降序
    h1 = ["序号", "行业", "类别", "岗位", "节点来源", "优先级", "样本数", "单位数", "证据强度",
          "原始岗位名（出现次数）", "学历分布", "专业分布", "职责原文（前3条）", "岗位卡阶段", "岗位卡缺口",
          "示例链接", "审查结论", "问题说明", "修正建议"]
    rows = []
    live = sorted((p for p, c in cards.items() if c["sample_n"] > 0), key=lambda p: -cards[p]["sample_n"])
    for i, path in enumerate(live, 1):
        c = cards[path]
        ps = by_node[path]
        parts = path.split(SEP)
        r = rc.get(path, {})
        rows.append([i, parts[0], parts[1] if len(parts) > 2 else "", parts[-1], c["source"], r.get("tier", ""),
                     c["sample_n"], len({p["company"] for p in ps}), None,
                     counts(Counter(p["role"] for p in ps)), counts(c["education"]), counts(c["majors"]),
                     "\n".join(c["duties"][:3]), r.get("completion_stage", "未建卡"),
                     "\n".join(r.get("missing", [])), (c.get("sample_urls") or [""])[0], "", "", ""])
    ws = wb.create_sheet("岗位审查")
    write_table(ws, h1, rows,
                {"序号": 6, "行业": 12, "类别": 12, "岗位": 20, "节点来源": 12, "优先级": 7, "样本数": 8, "单位数": 8,
                 "证据强度": 9, "原始岗位名（出现次数）": 40, "学历分布": 22, "专业分布": 28, "职责原文（前3条）": 50,
                 "岗位卡阶段": 16, "岗位卡缺口": 36, "示例链接": 30, "审查结论": 12, "问题说明": 30, "修正建议": 30},
                inputs=("审查结论", "问题说明", "修正建议"), verdicts=ROLE_VERDICTS)
    # 证据强度用公式，跟着单位数变
    for r in range(2, len(rows) + 2):
        ws[f"I{r}"] = f'=IF(H{r}<2,"证据少",IF(H{r}<5,"一般","较多"))'
    link(ws, "示例链接", h1, len(rows) + 1)
    n_role = len(rows)

    # 样本明细：全部去重后样本，含未挂上树的，便于逐条查归类
    h2 = ["样本ID", "挂到的岗位节点", "匹配方式", "原始岗位名", "单位", "专业类", "专业来源", "学历", "学历字段冲突",
          "发布日期", "信息来源", "链接", "归类判断", "应归到"]
    rows2 = sorted(([p["id"], p["node"] or "（未挂上）", p["node_match"], p["role"], p["company"],
                     "、".join(p["major_classes"]), p["major_source"], p["edu_level"], "是" if p["edu_conflict"] else "",
                     p["posted"], p["source"], p["url"], "", ""] for p in posts),
                   key=lambda r: (r[1] == "（未挂上）", r[1], r[3]))
    ws2 = wb.create_sheet("样本明细")
    write_table(ws2, h2, rows2,
                {"样本ID": 12, "挂到的岗位节点": 36, "匹配方式": 14, "原始岗位名": 26, "单位": 28, "专业类": 26,
                 "专业来源": 12, "学历": 8, "学历字段冲突": 8, "发布日期": 11, "信息来源": 14, "链接": 30,
                 "归类判断": 14, "应归到": 24},
                inputs=("归类判断", "应归到"), verdicts=SAMPLE_VERDICTS)
    link(ws2, "链接", h2, len(rows2) + 1)

    # 专业连线：专业类 × 岗位
    h3 = ["专业类", "行业", "岗位", "样本数", "单位数", "证据强度", "仅平台标签样本", "学历分布", "示例链接",
          "连线判断", "问题说明"]
    rows3 = []
    for e in sorted(edges, key=lambda e: (e["major_class"], -e["company_n"], -e["n"])):
        parts = e["node"].split(SEP)
        rows3.append([e["major_class"], parts[0], parts[-1], e["n"], e["company_n"], None,
                      e["platform_tag_only_n"], counts(e["education"]), (e["sample_urls"] or [""])[0], "", ""])
    ws3 = wb.create_sheet("专业连线")
    write_table(ws3, h3, rows3,
                {"专业类": 18, "行业": 14, "岗位": 24, "样本数": 8, "单位数": 8, "证据强度": 9, "仅平台标签样本": 10,
                 "学历分布": 26, "示例链接": 30, "连线判断": 12, "问题说明": 30},
                inputs=("连线判断", "问题说明"), verdicts=["合理", "不合理", "待定"])
    for r in range(2, len(rows3) + 2):
        ws3[f"F{r}"] = f'=IF(E{r}<2,"证据少",IF(E{r}<5,"一般","较多"))'
    link(ws3, "示例链接", h3, len(rows3) + 1)

    # 说明页：数据口径、填写方法、示例、审查进度（公式）
    lines = [
        ("职业库审查表（2026-10-07 导出）", None),
        ("", None),
        ("数据口径", None),
        (f"抓取样本 {stats['postings_total']} 条，去重后 {stats['single']} 条；挂上岗位树 {stats['mapped']} 条，未挂上 {stats['unmapped']} 条", None),
        ("来源：牛企直聘校招单岗页，每个关键词只取搜索第一页；数字是样本数，不是市场岗位数", None),
        ("单位数 = 挂到该岗位的不同招聘单位数；证据强度：1 家=证据少，2-4 家=一般，5 家及以上=较多", None),
        ("源文件：/Users/lucky_wang/Documents/Claude/projects/xiaogu-launch/data/careerdb/（cards.json、edges.json、postings_mapped.jsonl、role_cards/）", None),
        ("", None),
        ("怎么填", None),
        ("只填黄色列，其余列是库里导出的原始数据，改了不会回写", None),
        ("岗位审查：逐行看「原始岗位名」是否都属于这个岗位、职责是否对得上，在「审查结论」下拉里选", None),
        ("样本明细：岗位审查里发现归类问题时，按「挂到的岗位节点」筛选，逐条判断并写「应归到」", None),
        ("专业连线：判断「这个专业的人去做这个岗位」是否合理", None),
        ("", None),
        ("填写示例（岗位审查）", None),
        ("岗位=临床医生（科室未细分）｜审查结论=应拆分｜问题说明=原始岗位名里有病理医师、口腔医师等已写明科室的｜修正建议=写明科室的样本归到对应科室节点", None),
        ("", None),
        ("审查进度", None),
    ]
    for t, _ in lines:
        readme.append([t])
    readme.append(["岗位审查：已填结论", f'=COUNTA(岗位审查!Q2:Q{n_role + 1})', f"/ {n_role}"])
    base = readme.max_row
    for v in ROLE_VERDICTS:
        readme.append([f"  {v}", f'=COUNTIF(岗位审查!Q2:Q{n_role + 1},"{v}")'])
    readme.append(["样本明细：已判断", f'=COUNTA(样本明细!M2:M{len(rows2) + 1})', f"/ {len(rows2)}"])
    readme.append(["专业连线：已判断", f'=COUNTA(专业连线!J2:J{len(rows3) + 1})', f"/ {len(rows3)}"])
    readme.column_dimensions["A"].width = 110
    readme.column_dimensions["B"].width = 10
    for row in readme.iter_rows():
        for c in row:
            c.font = FONT
    for r in (1, 3, 9, 15, 18):
        readme.cell(row=r, column=1).font = Font(name="Arial", size=11 if r > 1 else 14, bold=True)
    readme.cell(row=base, column=1).font = FONT

    # 本机没有 LibreOffice 预算公式缓存值，让 Excel/Numbers 打开时整表重算
    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print(OUT, n_role, len(rows2), len(rows3))


if __name__ == "__main__":
    main()
